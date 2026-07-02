"""Vygeneruje Excel šablonu bez ukládání binárního .xlsx do gitu.

Použití:
    python tools/create_excel_template.py

Výstup:
    dist/MT5_Analysis_Portfolio_Dashboard.xlsx

Poznámka: Pro běh je potřeba balíček openpyxl (`pip install openpyxl`).
Samotná desktopová aplikace openpyxl nepotřebuje.
"""
from __future__ import annotations

from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError as exc:  # pragma: no cover - uživatelská chyba prostředí
    raise SystemExit("Chybí openpyxl. Nainstaluj ho příkazem: pip install openpyxl") from exc

OUTPUT = Path("dist/MT5_Analysis_Portfolio_Dashboard.xlsx")

README_LINES = [
    "MT5 Analysis & Portfolio Dashboard",
    "",
    "Použití:",
    "1. V listu CONFIG zkontrolujte MT5CommonFolderPath. Výchozí hodnota je %APPDATA%\\MetaQuotes\\Terminal\\Common\\Files\\.",
    "2. CSV soubory MT5_Analysis_Prices.csv a MT5_Portfolio_Positions.csv musí vytvářet váš MT5 Expert Advisor se středníkem jako oddělovačem.",
    "3. CSV soubory nenechávejte otevřené v Excelu – MT5 do zamčených souborů nemusí umět zapisovat.",
    "4. Analýzy zapisujte do listu ANALYZY. Verdict: BUY, AVOID, SHORT, WATCH. Status: OPEN, CLOSED, CANCELLED.",
    "5. Po importu VBA modulů spusťte RefreshAllMT5Data pro načtení CSV a přepočet dashboardu.",
    "6. Spusťte UpdateTrackedSymbols pro zápis otevřených symbolů z ANALYZY do TrackedSymbols.csv.",
    "",
    "Instalace VBA:",
    "Uložte tento sešit jako .xlsm, otevřete VBA editor (Alt+F11), importujte moduly ze složky vba/ a obsah vba/ThisWorkbook.cls vložte do objektu ThisWorkbook.",
]

CONFIG_ROWS = [
    ("MT5CommonFolderPath", "%APPDATA%\\MetaQuotes\\Terminal\\Common\\Files\\", "Složka Common\\Files MetaTraderu"),
    ("AnalysisPricesFile", "MT5_Analysis_Prices.csv", "CSV ceny"),
    ("PortfolioPositionsFile", "MT5_Portfolio_Positions.csv", "CSV portfolio"),
    ("TrackedSymbolsFile", "TrackedSymbols.csv", "CSV sledované symboly"),
    ("BuySuccessPct", 3, "BUY úspěch %"),
    ("BuyFailPct", -3, "BUY chyba %"),
    ("ShortSuccessPct", 3, "SHORT úspěch %"),
    ("ShortFailPct", -3, "SHORT chyba %"),
    ("AvoidBadPct", 5, "AVOID chyba při růstu %"),
    ("DefaultHorizonDays", 14, "Výchozí horizont"),
    ("AutoRefreshOnOpen", "FALSE", "Automatický import při otevření"),
]

TABLE_HEADERS = {
    "ANALYZY": ["AnalysisID", "AnalysisDateTime", "Symbol", "Verdict", "HorizonDays", "StartPrice", "TargetPct", "StopPct", "Risk", "Reason", "Status", "Notes"],
    "MT5_CENY": ["SnapshotTime", "Symbol", "Bid", "Ask", "Last", "Spread", "Point", "Digits", "PriceUsed"],
    "MT5_PORTFOLIO": ["SnapshotTime", "AccountLogin", "AccountCurrency", "Symbol", "Ticket", "Type", "Volume", "OpenTime", "OpenPrice", "CurrentPrice", "PercentFromOpen", "Profit", "Swap", "Magic", "Comment"],
    "VYSLEDKY_ANALYZ": ["AnalysisID", "AnalysisDateTime", "Symbol", "Verdict", "HorizonDays", "StartPrice", "CurrentPrice", "CurrentReturnPct", "Price1D", "Return1D", "Price5D", "Return5D", "Price14D", "Return14D", "HorizonPrice", "HorizonReturnPct", "MaxRunupPct", "MaxDrawdownPct", "Result", "Comment"],
    "PORTFOLIO": ["SnapshotTime", "Symbol", "Ticket", "Type", "Volume", "OpenTime", "OpenPrice", "CurrentPrice", "PercentFromOpen", "Profit", "Swap", "Comment"],
}

TABLE_NAMES = {
    "ANALYZY": "tblAnalyzy",
    "MT5_CENY": "tblMT5Ceny",
    "MT5_PORTFOLIO": "tblMT5Portfolio",
    "VYSLEDKY_ANALYZ": "tblVysledkyAnalyz",
    "PORTFOLIO": "tblPortfolioLatest",
}


def add_table(ws, name: str, rows: int, cols: int) -> None:
    ref = f"A1:{get_column_letter(cols)}{rows}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(table)


def build_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    for title in ["CONFIG", "ANALYZY", "MT5_CENY", "MT5_PORTFOLIO", "VYSLEDKY_ANALYZ", "PORTFOLIO", "DASHBOARD"]:
        wb.create_sheet(title)

    for row_number, line in enumerate(README_LINES, start=1):
        ws.cell(row_number, 1, line)
        ws.cell(row_number, 1).alignment = Alignment(wrap_text=True)
    ws["A1"].font = Font(size=16, bold=True)
    ws.column_dimensions["A"].width = 120

    cfg = wb["CONFIG"]
    cfg.append(["Key", "Value", "Popis"])
    for row in CONFIG_ROWS:
        cfg.append(row)
    cfg_table = Table(displayName="tblConfig", ref=f"A1:C{len(CONFIG_ROWS) + 1}")
    cfg_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    cfg.add_table(cfg_table)

    for sheet_name, headers in TABLE_HEADERS.items():
        sheet = wb[sheet_name]
        sheet.append(headers)
        sheet.append([""] * len(headers))
        add_table(sheet, TABLE_NAMES[sheet_name], 2, len(headers))
        for idx, header in enumerate(headers, start=1):
            sheet.column_dimensions[get_column_letter(idx)].width = max(12, min(24, len(header) + 4))

    ana = wb["ANALYZY"]
    verdict_validation = DataValidation(type="list", formula1='"BUY,AVOID,SHORT,WATCH"', allow_blank=False)
    status_validation = DataValidation(type="list", formula1='"OPEN,CLOSED,CANCELLED"', allow_blank=False)
    ana.add_data_validation(verdict_validation)
    ana.add_data_validation(status_validation)
    verdict_validation.add("D2:D10000")
    status_validation.add("K2:K10000")

    dash = wb["DASHBOARD"]
    dash["A1"] = "Dashboard analýz a MT5 portfolia"
    dash["A1"].font = Font(size=18, bold=True)
    labels = [
        "Celkem analýz", "Uzavřené/skórovatelné analýzy", "BUY hit rate", "SHORT hit rate", "AVOID hit rate",
        "HIT", "MISS", "NEUTRAL", "WATCH ONLY", "Aktuální profit MT5", "Počet otevřených pozic",
        "Nejlepší MT5 pozice", "Nejhorší MT5 pozice",
    ]
    for row_number, label in enumerate(labels, start=3):
        dash.cell(row_number, 1, label)
        dash.cell(row_number, 2, "doplní makro RefreshAllMT5Data")

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        sheet.freeze_panes = "A2"
    return wb


def main() -> None:
    OUTPUT.parent.mkdir(exist_ok=True)
    workbook = build_workbook()
    workbook.save(OUTPUT)
    print(f"Vytvořeno: {OUTPUT}")


if __name__ == "__main__":
    main()
