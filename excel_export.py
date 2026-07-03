"""Export dat aplikace do .xlsx."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from analysis_engine import Analysis, AnalysisResult
from dashboard import DashboardSummary
from mt5_data import PortfolioPosition, SymbolPortfolioRow


def export_to_excel(folder: Path, analyses: list[Analysis], results: list[AnalysisResult], current_positions: list[PortfolioPosition], history: list[PortfolioPosition], summary: DashboardSummary) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"MT5_Analysis_Portfolio_Export_{datetime.now():%Y%m%d_%H%M}.xlsx"
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(title: str, headers: list[str], rows: list[list[object]]):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        return ws

    sheet("Analýzy", ["AnalysisID", "AnalysisDateTime", "Symbol", "Verdict", "HorizonDays", "StartPrice", "TargetPct", "StopPct", "Risk", "Reason", "Status", "Notes"], [[a.analysis_id, a.analysis_datetime, a.symbol, a.verdict, a.horizon_days, a.start_price, a.target_pct, a.stop_pct, a.risk, a.reason, a.status, a.notes] for a in analyses])
    sheet("Výsledky analýz", ["AnalysisID", "Symbol", "Verdict", "CurrentReturnPct", "Return1D", "Return5D", "Return14D", "HorizonReturnPct", "MaxRunupPct", "MaxDrawdownPct", "Result"], [[r.analysis.analysis_id, r.analysis.symbol, r.analysis.verdict, r.current_return_pct, r.return_1d, r.return_5d, r.return_14d, r.horizon_return_pct, r.max_runup_pct, r.max_drawdown_pct, r.result] for r in results])
    portfolio_headers = ["SnapshotTime", "Symbol", "Ticket", "Type", "Volume", "OpenTime", "OpenPrice", "CurrentPrice", "PercentFromOpen", "Profit", "Swap", "Comment", "Exposure/Estimated"]
    portfolio_rows = lambda items: [[p.snapshot_time, p.symbol, p.ticket, p.type, p.volume, p.open_time, p.open_price, p.current_price, p.percent_from_open, p.profit, p.swap, p.comment, p.value_for_share] for p in items]
    sheet("Aktuální MT5 portfolio", portfolio_headers, portfolio_rows(current_positions))
    sheet("Historie MT5 portfolia", portfolio_headers, portfolio_rows(history))
    sheet("Dashboard summary", ["Metrika", "Hodnota"], [["Celkem analýz", summary.total_analyses], ["Skórovatelné analýzy", summary.scorable_analyses], ["BUY hit rate", summary.buy_hit_rate], ["SHORT hit rate", summary.short_hit_rate], ["AVOID hit rate", summary.avoid_hit_rate], ["Aktuální profit MT5", summary.portfolio_profit], ["Počet otevřených pozic", summary.open_positions], ["Celková hodnota/expozice", summary.total_value]])
    sheet("Podíl v portfoliu", ["Symbol", "Hodnota pozice", "Podíl %", "Profit", "WeightedPercentFromOpen", "Zdroj hodnoty"], [[r.symbol, r.value, r.share_pct, r.profit, r.weighted_percent_from_open, r.value_source] for r in summary.symbol_rows])
    sheet("Výkon akcií v %", ["Symbol", "Počet pozic", "WeightedPercentFromOpen", "Profit", "Exposure/Estimated", "Nejlepší ticket %", "Nejhorší ticket %"], [[r.symbol, r.positions_count, r.weighted_percent_from_open, r.profit, r.value, r.best_ticket_pct, r.worst_ticket_pct] for r in summary.symbol_rows])
    history_rows = []
    for symbol, points in summary.symbol_performance_history.items():
        for snapshot_time, percent in points:
            history_rows.append([snapshot_time, symbol, percent])
    sheet("Výkon akcií v čase", ["SnapshotTime", "Symbol", "PercentFromOpen"], history_rows)
    wb.save(path)
    return path
